import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import layout, palette

GRAY = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _rgb(hex_value):
    h = str(hex_value).lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _accent_rgb(accent):
    """The accent as *text* on white paper, so light accents stay legible.
    Background bands use the raw hex; only type goes through here."""
    return _rgb(palette.on_paper(palette.resolve(accent)["hex"]))


def _ink_rgb(accent_hex):
    """White or dark ink, whichever stays readable on this accent. Light
    accents would swallow white text, so the choice is measured, not assumed."""
    return _rgb(palette.ink_for(accent_hex))


def _shade(paragraph, hex_value):
    """Fill a paragraph's background - single-column shading, not a table, so
    it stays invisible to ATS text extraction while giving a colored band."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_value)
    paragraph._p.get_or_add_pPr().append(shd)


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")[:40] or "job"


def folder(base, job_id, company):
    path = Path(base) / (str(job_id).rjust(3, "0") + "-" + slug(company))
    path.mkdir(parents=True, exist_ok=True)
    return path


def ask_save_path(default_dir, default_filename):
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        chosen = filedialog.asksaveasfilename(
            parent=root,
            title="Save Word file",
            initialdir=str(default_dir),
            initialfile=default_filename,
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
        )
    finally:
        root.destroy()
    return Path(chosen) if chosen else None


FONT = layout.FONT_PT
SPACE = layout.SPACE_PT


def _base_document():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(FONT["body"])
    style.paragraph_format.space_after = Pt(SPACE["para_after"])
    return doc


def _run(para, text, bold=False, color=BLACK, size=FONT["bullet"]):
    run = para.add_run(text)
    run.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size)
    return run


def _space_after(doc):
    """Close a section so the next heading has room above it."""
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(SPACE["section_end"])


def _heading(doc, text, accent_hex, first=False):
    para = doc.add_paragraph()
    # No gap above the very first heading - it butts against the contact band.
    para.paragraph_format.space_before = Pt(0 if first else SPACE["section_before"])
    para.paragraph_format.space_after = Pt(SPACE["section_after"])
    _shade(para, accent_hex)
    _run(para, text, bold=True, color=_ink_rgb(accent_hex), size=FONT["heading"])
    return para


def cv_docx(content, path, accent=palette.DEFAULT_ACCENT):
    doc = _base_document()
    accent_rgb = _accent_rgb(accent)
    accent_hex = palette.resolve(accent)["hex"].lstrip("#")
    ink = _ink_rgb(accent_hex)
    identity = content.get("identity") or {}

    name_para = doc.add_paragraph()
    name_para.paragraph_format.space_before = Pt(6)
    name_para.paragraph_format.space_after = Pt(2)
    _shade(name_para, accent_hex)
    _run(name_para, identity.get("name", ""), bold=True, color=ink, size=FONT["name"])

    headline = identity.get("headline")
    if headline:
        headline_para = doc.add_paragraph()
        headline_para.paragraph_format.space_after = Pt(2)
        _shade(headline_para, accent_hex)
        _run(headline_para, headline, bold=True, color=ink, size=FONT["headline"])

    contact = " | ".join(
        str(v) for v in [
            identity.get("location"), identity.get("phone"),
            identity.get("email"), identity.get("linkedin"),
        ] if v
    )
    if contact:
        contact_para = doc.add_paragraph()
        contact_para.paragraph_format.space_after = Pt(6)
        _shade(contact_para, accent_hex)
        _run(contact_para, contact, bold=False, color=ink, size=FONT["contact"])

    summary = content.get("summary")
    if summary:
        _heading(doc, "PROFESSIONAL SUMMARY", accent_hex, first=True)
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(SPACE["section_end"])
        _run(para, summary, bold=False, color=BLACK, size=FONT["body"])

    for index, section in enumerate(content.get("sections") or []):
        kind = section["kind"]
        _heading(doc, section["heading"], accent_hex, first=not summary and index == 0)

        if kind == "skills":
            for tier in section["tiers"]:
                para = doc.add_paragraph()
                _run(para, tier["label"] + ": ", bold=True, color=accent_rgb,
                     size=FONT["body"])
                _run(para, ", ".join(tier["items"]), bold=False, color=BLACK,
                     size=FONT["body"])

        elif kind == "experience":
            for role_index, role in enumerate(section["roles"]):
                para = doc.add_paragraph()
                # Air above every employer but the first, so roles read as
                # separate blocks rather than one run-on list.
                para.paragraph_format.space_before = Pt(
                    0 if role_index == 0 else SPACE["role_before"]
                )
                para.paragraph_format.space_after = Pt(SPACE["role_after"])
                _run(para, role["title"], bold=True, color=BLACK, size=FONT["role"])
                _run(para, ", " + role["company"], bold=False, color=accent_rgb,
                     size=FONT["role"])
                if role.get("client"):
                    client_para = doc.add_paragraph()
                    client_para.paragraph_format.space_after = Pt(SPACE["role_after"])
                    _run(client_para, "Client: " + role["client"],
                         bold=False, color=accent_rgb, size=FONT["client"])
                loc_para = doc.add_paragraph()
                loc_para.paragraph_format.space_after = Pt(SPACE["meta_after"])
                location = role.get("location") or ""
                _run(loc_para, (location + "    " if location else "") + role["period"],
                     bold=False, color=GRAY, size=FONT["meta"])
                for bullet_index, bullet in enumerate(role["bullets"]):
                    para = doc.add_paragraph(style="List Bullet")
                    last = bullet_index == len(role["bullets"]) - 1
                    para.paragraph_format.space_after = Pt(
                        SPACE["bullets_after"] if last else SPACE["bullet_after"]
                    )
                    _run(para, bullet, bold=False, color=BLACK, size=FONT["bullet"])

        elif kind == "projects":
            for item in section["items"]:
                para = doc.add_paragraph(style="List Bullet")
                _run(para, item["name"] + ":", bold=True, color=BLACK, size=FONT["bullet"])
                _run(para, " " + item["text"], bold=False, color=BLACK,
                     size=FONT["bullet"])

        elif kind == "education":
            for item in section["items"]:
                para = doc.add_paragraph()
                _run(para, item["degree"], bold=True, color=BLACK, size=FONT["body"])
                _run(para, "  -  " + item["school"] + " (" + item["years"] + ")",
                     bold=False, color=BLACK, size=FONT["body"])

        elif kind == "certifications":
            for item in section["items"]:
                para = doc.add_paragraph(style="List Bullet")
                _run(para, item, bold=False, color=BLACK, size=FONT["bullet"])

        elif kind == "languages":
            for lang in section["items"]:
                text = (
                    (lang["name"] + " (" + lang["level"] + ")")
                    if lang.get("level") else lang["name"]
                )
                para = doc.add_paragraph(style="List Bullet")
                _run(para, text, bold=False, color=BLACK, size=FONT["bullet"])

        elif kind == "highlights":
            for item in section["items"]:
                para = doc.add_paragraph(style="List Bullet")
                _run(para, item, bold=False, color=BLACK, size=FONT["bullet"])

        # Every section closes the same way, so the next heading always has air
        # above it regardless of what the section ended with.
        _space_after(doc)

    doc.save(path)
    return path


def letter_docx(body, identity, path):
    doc = _base_document()

    name = doc.add_paragraph()
    run = name.add_run((identity or {}).get("name", ""))
    run.bold = True
    run.font.size = Pt(14)

    contact = " | ".join(
        str(v) for v in [
            (identity or {}).get("location"),
            (identity or {}).get("email"),
            (identity or {}).get("phone"),
        ] if v
    )
    if contact:
        doc.add_paragraph(contact)

    doc.add_paragraph("")
    doc.add_paragraph("Dear hiring team,")
    doc.add_paragraph("")

    for para in re.split(r"\n\s*\n", (body or "").strip()):
        cleaned = " ".join(para.split())
        if cleaned:
            doc.add_paragraph(cleaned)

    doc.add_paragraph("")
    doc.add_paragraph("Warm regards,")
    doc.add_paragraph((identity or {}).get("name", ""))

    doc.save(str(path))
    return path
